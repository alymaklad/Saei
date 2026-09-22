"""
Search tracing and the Excel debug report.

Also covers the recall property the trace exists to protect: title-mismatched
jobs must survive run_search so the semantic retriever can rescue them. That
was a real bug -- the first generated report showed zero "semantic" rows,
because run_search's title filter had already discarded every job the
semantic path exists to recover.
"""
import openpyxl
import pytest

import config
from agents import search_agent
from agents.search_trace import NullTrace, SearchTrace, new_trace


# ---- NullTrace: every method a no-op, never raises ---------------------------

def test_null_trace_absorbs_every_call():
    t = NullTrace()
    with t.stage("anything"):
        t.record_expansion(["a"], cached=False)
        t.record_source("s", "i", raw=1)
        t.record_retrieval({"title": "x"}, "token", "kept")
        t.record_ranking({"title": "x"}, {"score": 1, "breakdown": {}})
        t.record_gate({"title": "x"}, 0.5, True)
        t.record_error("s", "i", "boom")
        t.set_meta(anything=1)
    assert t.write() is None
    assert t.enabled is False


def test_new_trace_respects_the_config_switch(monkeypatch):
    monkeypatch.setattr(config, "SEARCH_DEBUG_REPORTS", False)
    assert isinstance(new_trace("x"), NullTrace)
    monkeypatch.setattr(config, "SEARCH_DEBUG_REPORTS", True)
    assert isinstance(new_trace("x"), SearchTrace)


def test_stage_timer_never_swallows_an_exception():
    """Tracing is observational -- it must not turn a real failure into a
    silent success."""
    t = SearchTrace("x")
    with pytest.raises(ValueError):
        with t.stage("boom"):
            raise ValueError("real failure")
    assert t.stages[0]["stage"] == "boom"  # still recorded


# ---- recording --------------------------------------------------------------

def test_expansion_marks_the_original_position():
    t = SearchTrace("AI Engineer")
    t.record_expansion(["AI Engineer", "ML Engineer"], cached=True)
    assert t.expansion[0]["is_original_position"] is True
    assert t.expansion[1]["is_original_position"] is False
    assert t.expansion[0]["origin"] == "cache"


def test_source_record_derives_the_drop_counts():
    t = SearchTrace()
    t.record_source("greenhouse", "acme", raw=100, after_filter=20, after_cap=5)
    row = t.sources[0]
    assert row["dropped_by_filter"] == 80
    assert row["dropped_by_cap"] == 15


def test_ranking_record_flattens_every_factor_into_columns():
    t = SearchTrace()
    result = {
        "score": 0.8,
        "breakdown": {"skills": {"score": 0.5, "weight": 0.32, "evidence": "half"}},
        "missing_skills": ["Rust"],
    }
    t.record_ranking({"title": "Dev"}, result)
    row = t.ranking[0]
    assert row["skills_score"] == 0.5
    assert row["skills_weight"] == 0.32
    assert row["skills_contribution"] == pytest.approx(0.16)
    assert row["missing_skills"] == "Rust"


# ---- workbook ---------------------------------------------------------------

def _trace_with_data():
    t = SearchTrace("Software Engineer")
    t.set_meta(seniority="", embedding_provider="ollama", match_threshold=0.0)
    t.record_expansion(["Software Engineer", "Backend Developer"], cached=False)
    t.record_source("greenhouse", "acme", raw=10, after_filter=3, after_cap=3, seconds=0.5)
    t.record_retrieval({"title": "A", "url": "u1"}, "token", "kept", "matched")
    t.record_retrieval({"title": "B", "url": "u2"}, "semantic", "kept", "close", semantic_score=0.9)
    t.record_retrieval({"title": "C", "url": "u3"}, "semantic", "dropped", "not close")
    t.record_ranking({"title": "A"}, {
        "score": 0.9,
        "breakdown": {"skills": {"score": 1.0, "weight": 0.32, "evidence": "all"}},
        "missing_skills": [],
    })
    t.record_gate({"title": "A", "match_score": 0.9}, 0.0, True)
    t.record_error("lever", "dead-slug", "404")
    t.stages.append({"stage": "retrieval", "seconds": 1.2})
    return t


def test_workbook_has_every_expected_sheet(tmp_path):
    path = _trace_with_data().write(str(tmp_path))
    assert path is not None
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == [
        "Summary", "Query Expansion", "Sources", "Retrieval",
        "Ranking", "Match Gate", "Source Errors", "Stage Timings",
    ]


def test_workbook_rows_land_on_the_right_sheets(tmp_path):
    path = _trace_with_data().write(str(tmp_path))
    wb = openpyxl.load_workbook(path)
    assert wb["Retrieval"].max_row == 4          # header + 3
    assert wb["Query Expansion"].max_row == 3    # header + 2
    assert wb["Source Errors"].max_row == 2      # header + 1


def test_summary_counts_are_formulas_not_baked_in_literals(tmp_path):
    """A hardcoded count goes stale the moment someone filters a detail
    sheet; a formula doesn't."""
    path = _trace_with_data().write(str(tmp_path))
    wb = openpyxl.load_workbook(path)
    ws = wb["Summary"]
    formulas = [
        c.value for row in ws.iter_rows(min_col=2, max_col=2)
        for c in row if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert len(formulas) >= 10
    assert any("COUNTIFS" in f for f in formulas)


def test_workbook_asks_excel_to_recalculate_on_open(tmp_path):
    """openpyxl writes formulas with no cached value, so without this every
    Summary cell would read blank until someone forced a recalc."""
    path = _trace_with_data().write(str(tmp_path))
    wb = openpyxl.load_workbook(path)
    assert wb.calculation.fullCalcOnLoad is True


def test_an_empty_run_still_produces_a_valid_workbook(tmp_path):
    """A search that found nothing is exactly when the report matters most."""
    path = SearchTrace("Nothing").write(str(tmp_path))
    assert path is not None
    wb = openpyxl.load_workbook(path)
    assert "Summary" in wb.sheetnames


def test_a_write_failure_never_propagates(monkeypatch, tmp_path):
    """A completed search must not be reported as failed because a debug
    artifact couldn't be saved."""
    t = _trace_with_data()
    monkeypatch.setattr(t, "_write", lambda d: (_ for _ in ()).throw(OSError("disk full")))
    assert t.write(str(tmp_path)) is None


def test_old_reports_are_pruned_to_the_keep_limit(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "SEARCH_REPORT_KEEP", 3)
    for _ in range(5):
        t = SearchTrace("x")
        # Distinct timestamps so filenames don't collide.
        import datetime
        t.started_at = t.started_at + datetime.timedelta(seconds=_)
        t.write(str(tmp_path))
    remaining = list(tmp_path.glob("search_report_*.xlsx"))
    assert len(remaining) == 3


def test_keep_of_zero_disables_pruning(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "SEARCH_REPORT_KEEP", 0)
    import datetime
    for i in range(4):
        t = SearchTrace("x")
        t.started_at = t.started_at + datetime.timedelta(seconds=i)
        t.write(str(tmp_path))
    assert len(list(tmp_path.glob("search_report_*.xlsx"))) == 4


# ---- the recall property the report exposed ---------------------------------

def _only_greenhouse(monkeypatch, board):
    monkeypatch.setattr(config, "GREENHOUSE_BOARD_TOKENS", ["acme"])
    monkeypatch.setattr(config, "LEVER_COMPANY_SLUGS", [])
    monkeypatch.setattr(search_agent, "search_greenhouse", lambda b: board)
    monkeypatch.setattr(search_agent, "search_remoteok", lambda *a, **k: [])
    monkeypatch.setattr(search_agent, "search_weworkremotely", lambda *a, **k: [])
    monkeypatch.setattr(search_agent, "search_from_watchlist", lambda *a, **k: [])
    monkeypatch.setattr(search_agent, "get_configured_sites", lambda: [])
    monkeypatch.setattr(search_agent, "search_serpapi", lambda q: [])


BOARD = [
    {"title": "Software Engineer", "url": "https://x/1", "description": "python"},
    {"title": "Backend Developer", "url": "https://x/2", "description": "python"},
]


def test_title_mismatches_are_discarded_by_default(monkeypatch):
    _only_greenhouse(monkeypatch, BOARD)
    jobs, _ = search_agent.run_search(position="Software Engineer",
                                      target_roles=["Software Engineer"])
    assert [j["title"] for j in jobs] == ["Software Engineer"]


def test_title_mismatches_are_carried_forward_when_requested(monkeypatch):
    """Without this the semantic retriever is dead weight -- run_search would
    already have thrown away every job it exists to rescue."""
    _only_greenhouse(monkeypatch, BOARD)
    jobs, _ = search_agent.run_search(position="Software Engineer",
                                      target_roles=["Software Engineer"],
                                      include_title_mismatches=True)
    assert {j["title"] for j in jobs} == {"Software Engineer", "Backend Developer"}


def test_seniority_mismatches_are_never_carried_forward(monkeypatch):
    """Seniority stays a hard filter even in rescue mode -- a 'senior' search
    shouldn't surface an internship via the semantic path."""
    board = [
        {"title": "Senior Software Engineer", "url": "https://x/1", "description": "d"},
        {"title": "Software Engineer Intern", "url": "https://x/2", "description": "d"},
    ]
    _only_greenhouse(monkeypatch, board)
    jobs, _ = search_agent.run_search(position="Software Engineer", seniority="senior",
                                      target_roles=["Software Engineer"],
                                      include_title_mismatches=True)
    assert [j["title"] for j in jobs] == ["Senior Software Engineer"]


def test_carried_candidates_are_capped(monkeypatch):
    """Embedding cost per run needs a hard ceiling -- a single Greenhouse
    board can return 500+ jobs."""
    monkeypatch.setattr(config, "SEMANTIC_CANDIDATE_CAP", 3)
    board = [{"title": f"Unrelated Role {i}", "url": f"https://x/{i}", "description": "d"}
             for i in range(50)]
    _only_greenhouse(monkeypatch, board)
    jobs, _ = search_agent.run_search(position="Software Engineer",
                                      target_roles=["Software Engineer"],
                                      include_title_mismatches=True)
    assert len(jobs) == 3
