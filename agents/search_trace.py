"""
Search pipeline tracing → multi-sheet Excel debug reports.

Records what happened at every stage of a search run and writes it to a
timestamped .xlsx so the pipeline can be audited after the fact: what the
Position expanded to, what each source actually returned before and after each
filter, which jobs the token and semantic paths kept or dropped and *why*,
every ranking factor's individual contribution to a job's match score, and
what the match-score gate held back.

Two design rules this module follows strictly:

  1. **Tracing never changes behavior.** Every record_* call is
     fire-and-forget, and a NullTrace (used when reporting is off, or when a
     caller passes nothing) makes each one a no-op. A failure while writing a
     report must never fail a search run -- if the report can't be written,
     the run has already done its real work.

  2. **Counts on the Summary sheet are formulas, not Python-computed
     literals**, so filtering or editing a detail sheet keeps the summary
     honest rather than leaving a stale number behind.
"""
import os
import time
from datetime import datetime, timezone

import config

# Column widths are set per sheet rather than autofit (openpyxl has no autofit),
# so the report is readable the moment it's opened.
_FONT = "Arial"


class NullTrace:
    """No-op stand-in used when reporting is disabled, so call sites never need
    an `if trace is not None` guard."""

    enabled = False

    def stage(self, *a, **k):
        return _NullTimer()

    def record_expansion(self, *a, **k): pass
    def record_source(self, *a, **k): pass
    def record_retrieval(self, *a, **k): pass
    def record_ranking(self, *a, **k): pass
    def record_gate(self, *a, **k): pass
    def record_error(self, *a, **k): pass
    def set_meta(self, *a, **k): pass
    def write(self, *a, **k): return None


class _NullTimer:
    def __enter__(self): return self
    def __exit__(self, *exc): return False


class _StageTimer:
    def __init__(self, trace, name):
        self.trace, self.name = trace, name

    def __enter__(self):
        self.started = time.perf_counter()
        return self

    def __exit__(self, *exc):
        self.trace.stages.append({
            "stage": self.name,
            "seconds": round(time.perf_counter() - self.started, 3),
        })
        return False  # never swallow an exception


class SearchTrace:
    """Collects per-stage records during one search run."""

    enabled = True

    def __init__(self, position: str = ""):
        self.started_at = datetime.now(timezone.utc)
        self.meta = {"position": position}
        self.stages: list[dict] = []
        self.expansion: list[dict] = []
        self.sources: list[dict] = []
        self.retrieval: list[dict] = []
        self.ranking: list[dict] = []
        self.gate: list[dict] = []
        self.errors: list[dict] = []

    # ---- recording ---------------------------------------------------------

    def stage(self, name: str):
        return _StageTimer(self, name)

    def set_meta(self, **kwargs):
        self.meta.update(kwargs)

    def record_expansion(self, roles: list[str], cached: bool, source: str = "llm"):
        for i, role in enumerate(roles):
            self.expansion.append({
                "n": i + 1,
                "role_phrase": role,
                "is_original_position": i == 0,
                "origin": "cache" if cached else source,
            })

    def record_source(self, source: str, identifier=None, role_phrase: str = "",
                      raw: int = 0, after_filter: int = 0, after_cap: int = 0,
                      seconds: float | None = None, error: str = ""):
        self.sources.append({
            "source": source,
            "identifier": identifier or "",
            "role_phrase": role_phrase,
            "raw_returned": raw,
            "after_title_filter": after_filter,
            "after_cap": after_cap,
            "dropped_by_filter": max(0, raw - after_filter),
            "dropped_by_cap": max(0, after_filter - after_cap),
            "seconds": round(seconds, 3) if seconds is not None else "",
            "error": error,
        })

    def record_retrieval(self, job: dict, path: str, outcome: str, reason: str = "",
                         semantic_score=None):
        self.retrieval.append({
            "source": job.get("source", ""),
            "title": job.get("title", ""),
            "company": job.get("company", ""),
            "url": job.get("url") or job.get("hostedUrl") or job.get("absolute_url") or "",
            "path": path,                 # token | semantic | role1
            "outcome": outcome,           # kept | dropped
            "semantic_score": round(semantic_score, 4) if isinstance(semantic_score, (int, float)) else "",
            "reason": reason,
            "description_chars": len(job.get("description") or ""),
        })

    def record_ranking(self, job: dict, result: dict):
        breakdown = result.get("breakdown") or {}
        row = {
            "title": job.get("title", ""),
            "company": job.get("company", ""),
            "source": job.get("source", ""),
            "url": job.get("url") or job.get("hostedUrl") or "",
            "match_score": result.get("score"),
        }
        # One column trio per factor: raw score, weight, weighted contribution.
        for name, detail in breakdown.items():
            row[f"{name}_score"] = detail.get("score")
            row[f"{name}_weight"] = detail.get("weight")
            row[f"{name}_contribution"] = round(
                (detail.get("score") or 0) * (detail.get("weight") or 0), 4
            )
            row[f"{name}_evidence"] = detail.get("evidence", "")
        row["missing_skills"] = ", ".join(result.get("missing_skills") or [])
        self.ranking.append(row)

    def record_gate(self, job: dict, threshold: float, passed: bool):
        self.gate.append({
            "title": job.get("title", ""),
            "company": job.get("company", ""),
            "url": job.get("url") or job.get("hostedUrl") or "",
            "match_score": job.get("match_score"),
            "threshold": threshold,
            "outcome": "processed" if passed else "held back",
        })

    def record_error(self, source: str, identifier, error: str):
        self.errors.append({
            "source": source,
            "identifier": identifier or "",
            "error": str(error),
        })

    # ---- writing -----------------------------------------------------------

    def write(self, directory: str | None = None) -> str | None:
        """Writes the workbook and returns its path (None if writing failed).

        Deliberately swallows its own exceptions: a search run that completed
        successfully must not be reported as failed because a debug artifact
        couldn't be saved.
        """
        try:
            return self._write(directory or config.SEARCH_REPORT_DIR)
        except Exception as exc:  # noqa: BLE001
            print(f"[search-trace] could not write the debug report: {exc}")
            return None

    def _write(self, directory: str) -> str:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        os.makedirs(directory, exist_ok=True)
        stamp = self.started_at.strftime("%Y%m%d_%H%M%S")
        path = os.path.join(directory, f"search_report_{stamp}.xlsx")

        wb = Workbook()
        # Formulas are written without cached values by openpyxl; this makes
        # Excel/LibreOffice compute them on open rather than showing blanks.
        wb.calculation.fullCalcOnLoad = True

        header_font = Font(name=_FONT, bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="44546A")
        body_font = Font(name=_FONT)

        def add_sheet(title: str, rows: list[dict], columns: list[str] | None = None,
                      widths: dict | None = None):
            ws = wb.create_sheet(title)
            columns = columns or (list(rows[0].keys()) if rows else ["(no rows)"])
            ws.append([c.replace("_", " ") for c in columns])
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for row in rows:
                ws.append([row.get(c, "") for c in columns])
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = body_font
            for i, col in enumerate(columns, start=1):
                letter = get_column_letter(i)
                ws.column_dimensions[letter].width = (widths or {}).get(col, min(max(len(col) + 4, 12), 48))
            ws.freeze_panes = "A2"
            if rows:
                ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
            return ws

        # --- detail sheets (created first so Summary can reference them) -----
        add_sheet("Query Expansion", self.expansion,
                  ["n", "role_phrase", "is_original_position", "origin"],
                  {"role_phrase": 42})
        add_sheet("Sources", self.sources,
                  ["source", "identifier", "role_phrase", "raw_returned",
                   "after_title_filter", "after_cap", "dropped_by_filter",
                   "dropped_by_cap", "seconds", "error"],
                  {"identifier": 34, "role_phrase": 28, "error": 60})
        add_sheet("Retrieval", self.retrieval,
                  ["source", "title", "company", "path", "outcome",
                   "semantic_score", "reason", "description_chars", "url"],
                  {"title": 46, "company": 24, "reason": 52, "url": 60})

        ranking_columns = []
        if self.ranking:
            base = ["title", "company", "source", "match_score"]
            factor_cols = [c for c in self.ranking[0] if c not in base and c not in ("url", "missing_skills")]
            ranking_columns = base + factor_cols + ["missing_skills", "url"]
        add_sheet("Ranking", self.ranking, ranking_columns or None,
                  {"title": 46, "company": 24, "missing_skills": 46, "url": 60})

        add_sheet("Match Gate", self.gate,
                  ["title", "company", "match_score", "threshold", "outcome", "url"],
                  {"title": 46, "company": 24, "url": 60})
        add_sheet("Source Errors", self.errors,
                  ["source", "identifier", "error"],
                  {"identifier": 40, "error": 80})
        add_sheet("Stage Timings", self.stages, ["stage", "seconds"], {"stage": 34})

        # --- Summary (first sheet) ------------------------------------------
        ws = wb.active
        ws.title = "Summary"
        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["B"].width = 54

        title_font = Font(name=_FONT, bold=True, size=14)
        section_font = Font(name=_FONT, bold=True)

        ws["A1"] = "Search run report"
        ws["A1"].font = title_font

        rows = [
            ("Run started (UTC)", self.started_at.strftime("%Y-%m-%d %H:%M:%S")),
            ("Position searched", self.meta.get("position", "")),
            ("Seniority filter", self.meta.get("seniority", "") or "(any)"),
            ("Max results per site", self.meta.get("max_results_per_site", "") or "(no cap)"),
            ("Max age (days)", self.meta.get("max_age_days", "") or "(no limit)"),
            (None, None),
            ("SETTINGS AT RUN TIME", None),
            ("Search-only mode",
             "YES -- orchestrator skipped, nothing saved to the database"
             if self.meta.get("search_only") else "no (full pipeline ran)"),
            ("Query expansion enabled", self.meta.get("query_expansion_enabled", "")),
            ("Embedding provider", self.meta.get("embedding_provider", "")),
            ("Semantic match threshold", self.meta.get("similarity_threshold", "")),
            ("Match score gate", self.meta.get("match_threshold", "")),
            ("SerpAPI phrase limit", self.meta.get("serpapi_expansion_limit", "")),
            (None, None),
            ("PIPELINE FUNNEL", None),
            ("Role phrases searched", "=COUNTA('Query Expansion'!B:B)-1"),
            ("Sources queried", "=COUNTA(Sources!A:A)-1"),
            ("Jobs returned by sources (raw)", "=SUM(Sources!D:D)"),
            ("Jobs surviving title filters", "=SUM(Sources!E:E)"),
            ("Jobs dropped by title filter", "=SUM(Sources!G:G)"),
            ("Jobs dropped by per-site cap", "=SUM(Sources!H:H)"),
            ("Retrieval rows recorded", "=COUNTA(Retrieval!A:A)-1"),
            ("  kept by token path", '=COUNTIFS(Retrieval!D:D,"token",Retrieval!E:E,"kept")'),
            ("  recovered by semantic path", '=COUNTIFS(Retrieval!D:D,"semantic",Retrieval!E:E,"kept")'),
            ("  dropped at retrieval", '=COUNTIF(Retrieval!E:E,"dropped")'),
            ("Candidates ranked", "=COUNTA(Ranking!A:A)-1"),
            ("Average match score", '=IFERROR(AVERAGE(Ranking!D:D),"n/a")'),
            ("Best match score", '=IFERROR(MAX(Ranking!D:D),"n/a")'),
            ("Processed (passed the gate)", '=COUNTIF(\'Match Gate\'!E:E,"processed")'),
            ("Held back by the gate", '=COUNTIF(\'Match Gate\'!E:E,"held back")'),
            ("Source errors", "=COUNTA('Source Errors'!A:A)-1"),
            ("Total pipeline seconds", "=SUM('Stage Timings'!B:B)"),
        ]

        r = 3
        for label, value in rows:
            if label is None:
                r += 1
                continue
            ws.cell(row=r, column=1, value=label).font = (
                section_font if value is None else body_font
            )
            if value is not None:
                cell = ws.cell(row=r, column=2, value=value)
                cell.font = body_font
            r += 1

        note = ws.cell(
            row=r + 1, column=1,
            value=("Every count above is a formula reading the detail sheets, so it stays "
                   "correct if you filter or edit them. Tracing is observational only — it "
                   "never changes what the search pipeline does."),
        )
        note.font = Font(name=_FONT, italic=True, size=9)
        note.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 3, end_column=2)

        wb.save(path)
        _prune_old_reports(directory, config.SEARCH_REPORT_KEEP)
        return path


def _prune_old_reports(directory: str, keep: int) -> None:
    """Keeps the newest `keep` reports. 0/negative disables pruning entirely."""
    if not keep or keep <= 0:
        return
    try:
        files = sorted(
            (f for f in os.listdir(directory)
             if f.startswith("search_report_") and f.endswith(".xlsx")),
            reverse=True,
        )
        for stale in files[keep:]:
            os.remove(os.path.join(directory, stale))
    except OSError:
        pass  # pruning is housekeeping -- never worth failing a run over


def new_trace(position: str = ""):
    """Returns a real trace when reporting is on, a no-op one otherwise."""
    return SearchTrace(position) if config.SEARCH_DEBUG_REPORTS else NullTrace()
